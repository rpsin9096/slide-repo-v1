#!/usr/bin/env python3
"""Construye el DOCX ilustrado a partir de content.py + figures/."""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from content import TITLE, SUBTITLE, INTRO, BLOCKS, FIGCAP, PERCUSSION_TABLE, APPENDIX_NOTES
from extract_figures import FIGURES

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIGDIR = ROOT / "figures"
OUT = ROOT / "Semiotecnia-del-Torax-y-Aparato-Respiratorio.docx"

NAVY = RGBColor(0x14, 0x3C, 0x50)
teal = RGBColor(0x0E, 0x7C, 0x7B)
GREY = RGBColor(0x5A, 0x6B, 0x76)

PROV = {fid: (pdfp, bookp) for fid, pdfp, bookp, _box, _t in FIGURES}
FIGREF = re.compile(r"(\(Fig\.[^)]*\))")


def fig_label(fid: str) -> str:
    if fid == "3-08a":
        return "Fig. 3-8A"
    if fid == "3-08b":
        return "Fig. 3-8B"
    return f"Fig. 3-{int(fid.split('-')[1])}"


def shade(cell, hexcolor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(cell, color="B9C6CC", sz="4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz); el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_runs_with_figrefs(par, text: str, size=10.5) -> None:
    for chunk in FIGREF.split(text):
        if not chunk:
            continue
        run = par.add_run(chunk)
        run.font.size = Pt(size)
        if chunk.startswith("(Fig."):
            run.bold = True
            run.font.color.rgb = teal


def styled_bullet(doc, term, tag, body, level):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.left_indent = Cm(0.7 + 0.5 * level)
    r = par.add_run(f"{term} ")
    r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = NAVY
    if tag:
        r2 = par.add_run(f"[{tag}] ")
        r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = teal
        r2.font.all_caps = True
    r3 = par.add_run("— " if not body.startswith((":", ";")) else "")
    r3.font.size = Pt(10.5)
    add_runs_with_figrefs(par, body)
    return par


def add_plate(doc, fids):
    n = len(fids)
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    widths = {1: Cm(9.2), 2: Cm(7.9), 3: Cm(5.2)}
    max_heights = {1: Cm(11.0), 2: Cm(9.8), 3: Cm(8.4)}
    from PIL import Image as PILImage
    for cell, fid in zip(table.rows[0].cells, fids):
        set_cell_borders(cell)
        cell.vertical_alignment = 1
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path = FIGDIR / f"fig-{fid}.jpg"
        with PILImage.open(img_path) as im:
            iw, ih = im.size
        target_w = widths[n] - Cm(0.35)
        target_h = target_w * ih / iw
        if target_h > max_heights[n]:
            target_h = max_heights[n]
            target_w = target_h * iw / ih
        p.add_run().add_picture(str(img_path), width=target_w, height=target_h)
        title, note = FIGCAP[fid]
        pdfp, bookp = PROV[fid]
        cap = cell.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        r = cap.add_run(f"{fig_label(fid)}. {title}.")
        r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = NAVY
        if note:
            rn = cap.add_run(f" {note}")
            rn.font.size = Pt(8)
        src = cell.add_paragraph()
        src.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = src.add_run(f"Recorte del PDF fuente · p. {pdfp} del PDF (p. {bookp} del libro) · fig-{fid}.jpg")
        rs.font.size = Pt(7); rs.font.color.rgb = GREY; rs.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_percussion_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, txt in enumerate(PERCUSSION_TABLE[0]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "143C50")
    for row in PERCUSSION_TABLE[1:]:
        cells = table.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(txt); r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True; r.font.color.rgb = NAVY
    # repetir encabezado
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("Tabla 3-1. Timbres percutorios: sustrato, topografía fisiológica y valor patológico. EIC: espacio intercostal.")
    r.font.size = Pt(8); r.font.color.rgb = GREY; r.italic = True


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    before = footer_p.add_run("Semiotecnia del tórax y aparato respiratorio  ·  ")
    before.font.size = Pt(7.5); before.font.color.rgb = GREY
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    footer_p._p.append(fld)
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header_p.add_run("Semiotecnia del tórax: aparato respiratorio · Figuras 3-1 a 3-35")
    hr.font.size = Pt(7.5); hr.font.color.rgb = GREY; hr.italic = True


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr in ("top_margin", "bottom_margin"):
        setattr(sec, attr, Cm(1.7))
    sec.left_margin = sec.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 16, 6), ("Heading 2", 12.5, teal, 12, 4), ("Title", 27, NAVY, 0, 6)]:
        st = styles[name]
        st.font.name = "Calibri"; st.font.size = Pt(size); st.font.color.rgb = color; st.font.bold = True
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    add_page_number_footer(doc)

    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CAPÍTULO 3 · SEMIOTECNIA — MANIOBRAS DE EXPLORACIÓN")
    r.font.size = Pt(9); r.font.color.rgb = teal; r.bold = True; r.font.all_caps = True
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run(SUBTITLE); rs.font.size = Pt(11.5); rs.font.color.rgb = GREY; rs.italic = True
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run().add_picture(str(FIGDIR / "fig-3-34.jpg"), width=Cm(7.2))
    capc = doc.add_paragraph()
    capc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = capc.add_run("Portada: Fig. 3-34 — distribución anterior de los ruidos respiratorios normales (detalle en cap. 4).")
    rc.font.size = Pt(7.5); rc.font.color.rgb = GREY; rc.italic = True

    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (num, name) in zip(t.rows[0].cells, [("1", "Inspección"), ("2", "Palpación"), ("3", "Percusión"), ("4", "Auscultación")]):
        shade(cell, "EAF3F3"); set_cell_borders(cell)
        pc = cell.paragraphs[0]; pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rn = pc.add_run(f"{num}\n"); rn.bold = True; rn.font.size = Pt(15); rn.font.color.rgb = teal
        rm = pc.add_run(name); rm.bold = True; rm.font.size = Pt(9.5); rm.font.color.rgb = NAVY

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(10)
    ri = intro.add_run(INTRO); ri.font.size = Pt(11)

    nota = doc.add_paragraph()
    rn = nota.add_run(
        "Nota de edición: el texto reproduce el material provisto; las 36 imágenes (Figs. 3-1 a 3-35, con 3-8 "
        "separada en A y B) son recortes de renders a 300 dpi del PDF adjunto, sin redibujo. Cada epígrafe indica "
        "la página del PDF y la página impresa del libro. La Fig. 3-23, presente en la fuente pero no citada en "
        "el texto base, se incluye en el Apéndice B. Detalle completo: Apéndice A y figuras-manifiesto.csv."
    )
    rn.font.size = Pt(8.5); rn.font.color.rgb = GREY; rn.italic = True

    doc.add_page_break()
    indice_p = doc.add_paragraph()
    indice_r = indice_p.add_run("Índice")
    indice_r.bold = True; indice_r.font.size = Pt(16); indice_r.font.color.rgb = NAVY; indice_r.font.name = "Calibri"
    indice_p.paragraph_format.space_after = Pt(6)
    toc_fields = doc.add_paragraph()
    run = toc_fields.add_run()
    fld_char = OxmlElement("w:fldChar"); fld_char.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char)
    run2 = toc_fields.add_run()
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    run2._r.append(instr)
    run3 = toc_fields.add_run()
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run3._r.append(fld_end)
    hint = doc.add_paragraph()
    rh = hint.add_run("Si el índice aparece vacío, haga clic derecho sobre él y elija «Actualizar campos». Contenido: 1. Inspección torácica · 2. Palpación sistemática · 3. Percusión: mapeo y timbres · 4. Auscultación · Apéndices A-B.")
    rh.font.size = Pt(8); rh.font.color.rgb = GREY; rh.italic = True

    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            doc.add_paragraph(block[1], style="Heading 1")
        elif kind == "h2":
            doc.add_paragraph(block[1], style="Heading 2")
        elif kind == "p":
            par = doc.add_paragraph()
            add_runs_with_figrefs(par, block[1])
        elif kind == "bullets":
            for term, tag, body, level in block[1]:
                styled_bullet(doc, term, tag, body, level)
        elif kind == "plate":
            add_plate(doc, block[1])
        elif kind == "perc_table":
            add_percussion_table(doc)

    # Apéndice A
    doc.add_page_break()
    doc.add_paragraph("Apéndice A. Índice de figuras y procedencia de los recortes", style="Heading 1")
    doc.add_paragraph(
        "Cada imagen es un recorte directo de un render a 300 dpi de la página indicada del PDF adjunto "
        "(«semiotecnia del torax.pdf», 14 páginas). La columna «Libro» da la página impresa del capítulo 3."
    )
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Figura", "Título (epígrafe original, resumido)", "PDF", "Libro", "Archivo"]
    for i, txt in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "143C50")
    for fid, pdfp, bookp, _box, orig_title in FIGURES:
        cells = table.add_row().cells
        vals = [fig_label(fid), orig_title, f"p. {pdfp}", f"p. {bookp}", f"fig-{fid}.jpg"]
        for i, val in enumerate(vals):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(7.5)
            if i == 0:
                r.bold = True; r.font.color.rgb = NAVY
    for note in APPENDIX_NOTES:
        par = doc.add_paragraph(style="List Bullet")
        r = par.add_run(note); r.font.size = Pt(9)

    # Apéndice B
    doc.add_paragraph("Apéndice B. Figura complementaria no citada en el texto base", style="Heading 1")
    doc.add_paragraph(
        "La Fig. 3-23 aparece en la fuente (p. 18 del libro) entre las maniobras de expansión y las vibraciones "
        "vocales, pero no es citada en el texto provisto. Se adjunta aquí para completar la serie 3-1 a 3-35 y "
        "como referencia topográfica previa al frémito táctil y la percusión anterior (2.º espacio intercostal)."
    )
    add_plate(doc, ["3-23"])

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Documento ilustrado con figuras 3-1 a 3-35 recortadas del PDF fuente"
    doc.core_properties.keywords = "semiotecnia, tórax, aparato respiratorio, inspección, palpación, percusión, auscultación"
    doc.save(OUT)
    print(f"OK -> {OUT} ({OUT.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
