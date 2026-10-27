# -*- coding: utf-8 -*-
"""Exporta o deck renderizado como um SVG por slide (helper de verificação).

O oraculo `--audit-safe-zone` do FocusGate le viewBox de arquivos .svg: cada
slide vira um SVG 16:9 com as caixas reais desenhadas — o oraculo externo
valida o lienzo do pptx RENDERIZADO, nao a constante do codigo.
"""

import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from studydeck import manifest as MF           # noqa: E402
from studydeck import render as RD             # noqa: E402

SCALE = 1000          # 1 in -> 1000 unidades (viewBox 0 0 13333 7500)
IN = 914400           # EMU por polegada


def _inches(emu):
    return (emu or 0) / IN


def _svg_for_slide(shapes, width_u, height_u):
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" '
             f'viewBox="0 0 {width_u} {height_u}" '
             f'width="{width_u}" height="{height_u}">']
    parts.append(f'<rect x="0" y="0" width="{width_u}" height="{height_u}" '
                 f'fill="#F8F9FF"/>')
    for shape in shapes:
        try:
            if shape.left is None:
                continue
            x = round(_inches(shape.left) * SCALE, 1)
            y = round(_inches(shape.top) * SCALE, 1)
            w = round(_inches(shape.width) * SCALE, 1)
            h = round(_inches(shape.height) * SCALE, 1)
        except (TypeError, ValueError):
            continue
        parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" '
                     f'fill="none" stroke="#3854CC" stroke-width="8"/>')
    parts.append('</svg>')
    return '\n'.join(parts)


def main(argv):
    if len(argv) < 3:
        print('Uso: export_deck_svg.py <manifest.py> <out_dir>', file=sys.stderr)
        return 2
    manifest_path, out_dir = argv[1], argv[2]
    pathlib.Path(out_dir).mkdir(parents=True, exist_ok=True)
    deck = MF.load(manifest_path)
    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    tmp.close()
    try:
        _n, fails = RD.render(deck, tmp.name)
        if fails:
            print(f'EXPORT_FAIL: deck com violacao de contraste: {fails[:2]}',
                  file=sys.stderr)
            return 1
        from pptx import Presentation
        prs = Presentation(tmp.name)
        width_u = round(_inches(prs.slide_width) * SCALE)
        height_u = round(_inches(prs.slide_height) * SCALE)
        written = 0
        for i, slide in enumerate(prs.slides, 1):
            out = pathlib.Path(out_dir) / f'slide_{i:02d}.svg'
            out.write_text(_svg_for_slide(slide.shapes, width_u, height_u),
                           encoding='utf-8')
            written += 1
        print(f'EXPORT_OK: {written} slide(s) -> {out_dir} '
              f'(viewBox {width_u}x{height_u})')
        return 0
    finally:
        pathlib.Path(tmp.name).unlink(missing_ok=True)


if __name__ == '__main__':
    sys.exit(main(sys.argv))