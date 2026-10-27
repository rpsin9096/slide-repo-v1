# -*- coding: utf-8 -*-
"""Pipeline de render: monta o .pptx y falla CERRADO si un par texto/superficie
queda bajo el umbral WCAG. Un deck con contraste insuficiente no se entrega."""

import pathlib
import tempfile
import textwrap
import unittest
from unittest import mock

from pptx import Presentation

from studydeck import manifest as M
from studydeck import render as R

SLIDES = [
    {'module': 'title', 'title': 'Inflamação aguda', 'subtitle': 'Do estímulo à resolução',
     'course': 'Fisiopatologia', 'institution': 'UCP'},
    {'module': 'definition', 'kicker': 'Conceito', 'title': 'Definição operativa',
     'main_statement': 'Resposta vascular estereotipada do tecido vivo diante da lesao.',
     'key_points': [{'title': 'Vascular', 'desc': 'Vasodilatacao e permeabilidade.'},
                    {'title': 'Celular', 'desc': 'Neutrofilos e monocitos.'}]},
    {'module': 'section_divider', 'section_title': 'A resposta inflamatória',
     'section_subtitle': 'Vasodilatacao, permeabilidade e migracao leucocitaria.'},
]


def write_manifest(tmp, slides=SLIDES, extra=''):
    body = ',\n'.join(repr(s) for s in slides)
    src = (f"COURSE_INFO = {{'course': 'Fisiopatologia', 'institution': 'UCP'}}\n"
           f"CONFIG = {{'LANG': 'pt-BR', 'PROFILE': 'short'}}{extra}\n"
           f"SLIDES = [{body}]\n")
    p = pathlib.Path(tmp) / 'm.py'
    p.write_text(src, encoding='utf-8')
    return str(p)


class TestRender(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = self._tmp.name
        self.out = f'{self.tmp}/deck.pptx'

    def tearDown(self):
        self._tmp.cleanup()

    def test_renders_every_slide(self):
        doc = M.load(write_manifest(self.tmp))
        n, fails = R.render(doc, self.out)
        self.assertEqual(n, 3)
        self.assertEqual(fails, [])

    def test_writes_a_readable_pptx(self):
        doc = M.load(write_manifest(self.tmp))
        R.render(doc, self.out)
        prs = Presentation(self.out)
        self.assertEqual(len(prs.slides.__iter__.__self__._sldIdLst), 3)

    def test_canvas_is_16_9(self):
        doc = M.load(write_manifest(self.tmp))
        R.render(doc, self.out)
        prs = Presentation(self.out)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)

    def test_uses_montserrat(self):
        doc = M.load(write_manifest(self.tmp))
        R.render(doc, self.out)
        prs = Presentation(self.out)
        fonts = {r.font.name for sl in prs.slides for sh in sl.shapes
                 if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
        self.assertEqual(fonts - {None}, {'Montserrat'})

    def test_contrast_violation_deletes_the_output(self):
        doc = M.load(write_manifest(self.tmp))
        original = R.MODULES['definition']['render']

        def saboteur(slide, d, pal, info):
            original(slide, d, pal, info)
            from studydeck.primitives import tx
            tx(slide, 1.0, 5.0, 4.0, 0.4, 'texto ilegível', size=12, color='C8C8FF')

        with mock.patch.dict(R.MODULES['definition'], {'render': saboteur}):
            n, fails = R.render(doc, self.out)
        self.assertTrue(fails)
        self.assertFalse(pathlib.Path(self.out).exists())

    def test_web_source_renders_a_labelled_panel_without_downloading(self):
        slides = [dict(SLIDES[0]),
                  {'module': 'figure', 'kicker': 'Evidencia', 'title': 'Serologia',
                   'image': 'https://exemplo.org/nunca-baixado.jpg',
                   'caption': 'Anti-CCP elevado em paciente seronegativo para fator reumatoide.',
                   'credit': 'Arquivo proprio'}]
        doc = M.load(write_manifest(self.tmp, slides))
        n, fails = R.render(doc, self.out)
        self.assertEqual(fails, [])
        self.assertFalse(pathlib.Path('nunca-baixado.jpg').exists())

    def test_unknown_module_is_rejected(self):
        doc = M.load(write_manifest(self.tmp, [{'module': 'inventado'}]))
        with self.assertRaises(KeyError):
            R.render(doc, self.out)

    def test_render_is_idempotent_on_output(self):
        doc = M.load(write_manifest(self.tmp))
        R.render(doc, self.out)
        n2, fails2 = R.render(doc, self.out)
        self.assertEqual((n2, fails2), (3, []))

    def test_missing_local_image_raises(self):
        slides = [dict(SLIDES[0]),
                  {'module': 'figure', 'title': 'Figura', 'image': f'{self.tmp}/ausente.png',
                   'caption': 'Corte coronal do lobo temporal com atrofia medial.',
                   'credit': 'Arquivo proprio'}]
        doc = M.load(write_manifest(self.tmp, slides))
        with self.assertRaises(FileNotFoundError):
            R.render(doc, self.out)


if __name__ == "__main__":
    unittest.main()