from pathlib import Path
import re, io, json
import pymupdf as fitz
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RImage, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'; AS=OUT/'figuras'; AS.mkdir(exist_ok=True)
src=fitz.open('/tmp/arena-coding-uploads/semiotecnia del torax.pdf')
# Crop rectangles are fractions of rendered page width and height, excluding original captions.
boxes={
1:(3,.08,.63,.52,.83),2:(3,.53,.63,.95,.84),
3:(4,.08,.13,.52,.35),4:(4,.54,.12,.96,.36),5:(4,.09,.58,.51,.79),6:(4,.57,.49,.92,.82),
7:(5,.31,.18,.66,.53),8:(5,.08,.59,.86,.87),
9:(6,.09,.11,.51,.43),10:(6,.53,.10,.95,.43),11:(6,.16,.52,.48,.85),12:(6,.54,.51,.95,.85),
13:(7,.08,.18,.50,.49),14:(7,.53,.18,.96,.50),15:(7,.12,.57,.50,.90),16:(7,.53,.58,.96,.90),
17:(8,.10,.10,.51,.46),18:(8,.55,.10,.96,.46),19:(8,.10,.56,.54,.83),
20:(9,.49,.20,.92,.50),21:(9,.07,.57,.51,.87),22:(9,.54,.59,.91,.90),
24:(10,.53,.10,.95,.34),25:(10,.16,.56,.55,.89),
26:(11,.10,.24,.51,.49),27:(11,.55,.22,.95,.49),28:(11,.55,.65,.95,.94),
29:(12,.10,.11,.52,.37),30:(12,.55,.12,.96,.29),31:(12,.12,.46,.52,.78),32:(12,.59,.46,.94,.77),
33:(13,.07,.25,.52,.55),34:(13,.53,.25,.94,.55),35:(13,.53,.66,.94,.96)}
names={1:'Tórax enfisematoso',2:'Tórax en quilla',3:'Tórax excavado',4:'Tórax piriforme',5:'Tórax paralítico o asténico',6:'Cifosis senil',7:'Cifosis juvenil',8:'Escoliosis: A, sinistroconvexa; B, dextroconvexa',9:'Mano de escultor',10:'Evaluación del trofismo',11:'Evaluación del tono muscular',12:'Palpación axilar: grupo central',13:'Ganglios supraclaviculares',14:'Palpación axilar: grupo externo',15:'Palpación axilar: grupo pectoral',16:'Palpación axilar: grupo subescapular',17:'Palpación infraclavicular',18:'Resistencia torácica',19:'Expansión apical: maniobra de Ruault',20:'Expansión infraclavicular',21:'Expansión de bases anteriores',22:'Expansión de bases posteriores',24:'Vibraciones vocales: cara anterior',25:'Vibraciones vocales: cara posterior',26:'Vibraciones vocales: delimitación basal',27:'Percusión de campos de Krönig',28:'Percusión directa de clavículas',29:'Percusión del esternón',30:'Percusión anterior',31:'Percusión posterior',32:'Excursión de las bases pulmonares',33:'Percusión de la columna',34:'Auscultación anterior',35:'Auscultación posterior'}
manifest=[]
for n,(p,x0,y0,x1,y1) in boxes.items():
 page=src[p-1]; w,h=page.rect.width,page.rect.height
 # Render the full page first; crops are taken from the raster, not embedded images.
 pix=page.get_pixmap(matrix=fitz.Matrix(3,3)); im=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
 crop=im.crop((int(x0*im.width),int(y0*im.height),int(x1*im.width),int(y1*im.height)))
 crop.save(AS/f'fig-3-{n:02}.png')
 manifest.append({'figure':f'3-{n}','pdf_page':p,'printed_page':p+8,'crop_fraction':[x0,y0,x1,y1],'caption':names[n]})
(AS/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
font='/usr/share/fonts/truetype/dejavu/'
for name,file in [('DejaVu','DejaVuSans.ttf'),('DejaVu-Bold','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,font+file))
pdfmetrics.registerFontFamily('DejaVu',normal='DejaVu',bold='DejaVu-Bold',italic='DejaVu',boldItalic='DejaVu-Bold')
styles=getSampleStyleSheet()
for key in ['Normal','BodyText','Title','Heading1','Heading2']:
 styles[key].fontName='DejaVu'; styles[key].textColor=colors.HexColor('#203747')
styles['Normal'].fontSize=9.4;styles['Normal'].leading=14;styles['Normal'].spaceAfter=7
styles['Title'].fontSize=28;styles['Title'].leading=35
styles['Heading1'].fontSize=21;styles['Heading1'].leading=27;styles['Heading1'].spaceAfter=15
styles['Heading2'].fontSize=12;styles['Heading2'].leading=17;styles['Heading2'].spaceBefore=12
cap=ParagraphStyle('caption',parent=styles['Normal'],fontSize=8,leading=11,spaceBefore=6)
small=ParagraphStyle('small',parent=styles['Normal'],fontSize=7.5,leading=10)
doc=Document(); sec=doc.sections[0];sec.header.paragraphs[0].text='GUÍA ILUSTRADA  /  SEMIOLOGÍA RESPIRATORIA'
sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
for s in ['Title','Heading 1','Heading 2']:doc.styles[s].font.color.rgb=RGBColor.from_string('203747')
sec.footer.paragraphs[0].text='Semiotecnia del tórax • Material de estudio • Figuras del PDF facilitado'
story=[]
def para(t,sty='Normal'):
 return Paragraph(re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',t),styles[sty])
def dpara(t):
 p=doc.add_paragraph()
 for i,s in enumerate(re.split(r'\*\*(.*?)\*\*',t)):p.add_run(s).bold=bool(i%2)

def figures(nums):
 for off in range(0,len(nums),4):
  chunk=nums[off:off+4];story.append(PageBreak());story.append(para('Atlas de maniobras y hallazgos','Heading2'));doc.add_page_break();doc.add_heading('Atlas de maniobras y hallazgos',2)
  rows=[]; dt=doc.add_table(rows=0,cols=2)
  for i in range(0,len(chunk),2):
   row=[]; cells=dt.add_row().cells
   for j,n in enumerate(chunk[i:i+2]):
    path=AS/f'fig-3-{n:02}.png';im=Image.open(path);w,h=im.size;scale=min(220/w,225/h)
    caption=f'Fig. 3-{n}. {names[n]}.'
    row.append([RImage(str(path),w*scale,h*scale),Paragraph(caption,cap)])
    p=cells[j].paragraphs[0];p.add_run().add_picture(str(path),width=Inches(w*min(2.75/w,2.8/h)),height=Inches(h*min(2.75/w,2.8/h)));cells[j].add_paragraph(caption,style='Caption')
   if len(row)==1:row.append('')
   rows.append(row)
  t=Table(rows,colWidths=[243,243],hAlign='LEFT');t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),20)]));story.append(t)
 story.append(PageBreak());doc.add_page_break()
rows=[['Sonido','Sustrato biofísico','Topografía fisiológica','Hallazgo patológico'],['Sonoridad','Resonancia del parénquima alveolar aireado.','Campos de Krönig; dorso hasta bases; derecho hasta 4.º EIC e izquierdo hasta 3.º EIC.','Atenuada en atelectasias o consolidaciones.'],['Hipersonoridad','Tono más grave e intensidad aumentada por exceso aéreo.','No existe en reposo fisiológico.','Enfisema, crisis asmática, neumotórax.'],['Timpanismo','Sonido musical agudo en cavidad aérea tensa cerrada.','Espacio semilunar de Traube (cámara gástrica).','Neumotórax a tensión; cavernas superficiales de pared lisa.'],['Matidez','Golpe sordo y breve sobre masa sólida sin gas.','Desde 5.º EIC derecho (hígado); área precordial cardíaca.','Consolidación neumónica masiva; derrame pleural.'],['Submatidez','Transición entre aire alveolar y solidez visceral.','4.º a 5.º EIC derecho; área periprecordial.','Derrames laminares incipientes; condensaciones incompletas.']]
for line in (OUT/'semiotecnia-del-torax.md').read_text().splitlines():
 if not line:continue
 if line.startswith('@fig '):figures([int(x) for x in line[5:].split(',')]);continue
 if line=='@table':
  t=Table([[Paragraph(v,small) for v in r] for r in rows],colWidths=[75,135,143,133],repeatRows=1)
  t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbe9ed')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#b7c9d0')),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]));story.append(t)
  dt=doc.add_table(rows=1,cols=4);dt.style='Light Shading Accent 1'
  for c,v in zip(dt.rows[0].cells,rows[0]):c.text=v
  for r in rows[1:]:
   for c,v in zip(dt.add_row().cells,r):c.text=v
  continue
 if line.startswith('# '):story.extend([Spacer(1,90),para(line[2:],'Title')]);doc.add_heading(line[2:],0)
 elif line.startswith('## '):
  story.append(PageBreak());story.append(para(line[3:],'Heading1'));doc.add_page_break();doc.add_heading(line[3:],1)
 elif line.startswith('### '):story.append(para(line[4:],'Heading2'));doc.add_heading(line[4:],2)
 else:story.append(para(line));dpara(line)
def footer(canvas,doc):
 canvas.setFont('DejaVu',7);canvas.setFillColor(colors.HexColor('#627781'));canvas.drawString(54,28,'SEMIOLOGÍA RESPIRATORIA • GUÍA ILUSTRADA');canvas.drawRightString(541,28,str(doc.page))
story = [item for i,item in enumerate(story) if not (isinstance(item,PageBreak) and i and isinstance(story[i-1],PageBreak))]
# Remove consecutive explicit page breaks from the editable version as well.
from docx.oxml.ns import qn
previous_break=False
for element in list(doc.element.body):
    is_break = element.tag == qn('w:p') and bool(element.xpath('.//w:br[@w:type="page"]')) and not ''.join(element.itertext()).strip()
    if is_break and previous_break: element.getparent().remove(element)
    else: previous_break=is_break
SimpleDocTemplate(str(OUT/'semiotecnia-del-torax.pdf'),pagesize=(595.28,841.89),rightMargin=54,leftMargin=54,topMargin=48,bottomMargin=48,title='Semiotecnia del tórax y aparato respiratorio',author='Material facilitado por el usuario').build(story,onFirstPage=footer,onLaterPages=footer)
doc.save(OUT/'semiotecnia-del-torax.docx')
print('Created PDF, DOCX and',len(boxes),'cropped figures')
