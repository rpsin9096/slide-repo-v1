# -*- coding: utf-8 -*-
"""Renderers tier scholarly: cita, tutoria, debate, evidencia PICO, viva voce, mnemotecnia, etimologia, conjuntos, grafico, referencias y autoridad."""

import math
import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..contrast import ratio, shade_until, WCAG_AA_TEXT
from ..geometry import (
    CANVAS_H, CANVAS_W, CONTENT_W, SAFE_X0, SAFE_Y0,
    row_layout, stack_layout,
)
from ..i18n import T, bloom
from ..primitives import (
    C, bg, footnote, note_surface, para, pill, place_image, rect,
    set_fill_alpha, textframe, tx, unmeasured,
)
from ..typography import FONT_BODY, FONT_TITLE

def m_quote_block(slide, d, pal, info):
    """Epigraph: seminal quotation in scholarly serif with attribution."""
    tx(slide, 1.0, 1.45, 1.15, 1.3, '\u201C', size=88, color=pal["accent"], bold=True,
       font=FONT_TITLE, deco=True)  # ornamento exento del gate
    tx(slide, 2.2, 1.9, 9.4, 2.3, d.get("quote", ""), size=20, color=pal["text"], italic=True,
       font=FONT_TITLE, align=PP_ALIGN.CENTER, line_sp=1.25)
    rect(slide, 5.87, 4.45, 1.6, 0.03, fill=pal["accent"], rounded=False)
    tx(slide, 2.2, 4.65, 9.4, 0.35, "— " + d.get("author", "").upper(), size=13, color=pal["primary"],
       bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE)
    src = " · ".join(x for x in (d.get("source", ""), d.get("year", "")) if x)
    if src:
        tx(slide, 2.2, 5.02, 9.4, 0.3, src, size=10.5, color=pal["text_muted"],
           italic=True, align=PP_ALIGN.CENTER)
    if d.get("context"):
        rect(slide, 2.2, 5.55, 9.4, 0.8, fill=pal["soft"], line=pal["border"], lw=0.4, rounded=True)
        tx(slide, 2.45, 5.66, 8.9, 0.6, d["context"], size=10.5, color=pal["text"],
           align=PP_ALIGN.CENTER, line_sp=1.1)
    footnote(slide, pal, d)

def m_socratic_question(slide, d, pal, info):
    """Tutorial prompt: hero question + numbered reasoning scaffold."""
    ink = pal.ink_accent()
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), Inches(1.8), Inches(1.5), Inches(1.5))
    ov.fill.solid(); ov.fill.fore_color.rgb = C(pal["primary"])
    note_surface(pal["primary"], 0.75, 1.8, 1.5, 1.5)
    ov.line.color.rgb = C(pal["accent"]); ov.line.width = Pt(1.5)
    ov.shadow.inherit = False
    otf = ov.text_frame
    otf.word_wrap = False
    otf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(otf, "?", 44, pal.tokens['on_primary'], bold=True, font=FONT_TITLE, first=True, align=PP_ALIGN.CENTER, sp_after=0)
    tx(slide, 2.55, 1.75, 9.9, 0.3, T('tutorial'), size=10, color=ink, bold=True)
    tx(slide, 2.55, 2.1, 9.9, 1.6, d.get("question", ""), size=20, color=pal["primary"],
       italic=True, font=FONT_TITLE, line_sp=1.2)
    if d.get("context"):
        tx(slide, 2.55, 3.75, 9.9, 0.45, d["context"], size=10.5,
           color=pal["text_muted"], line_sp=1.15)
    sc = d.get("scaffold", [])
    n = len(sc)
    if n:
        gap = 0.25
        w, sx = row_layout(SAFE_X0, CONTENT_W, n, gap)
        for i, s in enumerate(sc):
            x = sx + i * (w + gap)
            rect(slide, x, 4.35, w, 1.75, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
            rect(slide, x, 4.35, w, 0.05, fill=pal["accent"], rounded=False)
            bd = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(4.55), Inches(0.4), Inches(0.4))
            bd.fill.solid(); bd.fill.fore_color.rgb = C(pal["primary"])
            note_surface(pal["primary"], x + 0.18, 4.55, 0.4, 0.4)
            bd.line.fill.background(); bd.shadow.inherit = False
            btf = bd.text_frame
            btf.word_wrap = False
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(btf, str(i + 1), 12, pal.tokens['on_primary'], bold=True, first=True, align=PP_ALIGN.CENTER, sp_after=0)
            tf = textframe(slide, x + 0.7, 4.55, w - 0.9, 1.4)
            para(tf, s, 10.8, pal["text"], first=True, line_sp=1.12)
    if d.get("reading"):
        tx(slide, SAFE_X0, 6.30, CONTENT_W, 0.3, d["reading"], size=9.5, color=pal["text_muted"], italic=True)
    footnote(slide, pal, d)

def m_oxford_union(slide, d, pal, info):
    """Union-style motion debate: motion banner, proposition vs opposition, verdict."""
    rect(slide, SAFE_X0, 1.5, CONTENT_W, 1.15, fill=pal["primary"], line=None, rounded=True)
    tx(slide, 0.95, 1.62, 11.4, 0.28, T('house'), size=9.5, color=pal.tokens['on_primary'], bold=True)
    tx(slide, 0.95, 1.92, 11.4, 0.68, '"' + d.get("motion", "") + '"', size=17, color=pal.tokens['on_primary'],
       bold=True, font=FONT_TITLE, line_sp=1.05)
    for x, key, bar in ((SAFE_X0, 'pro_card', pal["primary"]), (6.77, 'con_card', pal["accent"])):
        card = d.get(key, {})
        rect(slide, x, 2.9, 5.88, 2.95, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, x, 2.9, 5.88, 0.44, fill=bar, rounded=False)
        tx(slide, x + 0.25, 2.97, 5.38, 0.32, card.get("title", "").upper(), size=10.5,
           color=pal.on(bar), bold=True)
        tf = textframe(slide, x + 0.25, 3.5, 5.38, 2.2)
        for i, p in enumerate(card.get("points", [])):
            para(tf, "• " + p, 11.0, pal["text"], first=(i == 0), sp_after=6, line_sp=1.1)
    rect(slide, SAFE_X0, 6.05, CONTENT_W, 0.55, fill=pal["soft"], line=pal["accent"], lw=0.9, rounded=True)
    tx(slide, 0.9, 6.09, 11.5, 0.48, T('verdict') + " " + d.get("verdict", ""), size=11,
       color=pal["primary"], bold=True, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
    footnote(slide, pal, d, y=6.66)

def m_evidence_card(slide, d, pal, info):
    """Landmark trial appraisal: PICO rows + effect hero + limitation strip."""
    rect(slide, SAFE_X0, 1.5, CONTENT_W, 0.95, fill=pal["primary"], line=None, rounded=True)
    tx(slide, 0.95, 1.62, 9.0, 0.42, d.get("study_name", ""), size=17, color=pal.tokens['on_primary'],
       bold=True, font=FONT_TITLE)
    if d.get("year"):
        rect(slide, 10.85, 1.64, 1.5, 0.38, fill=pal["accent"], line=None, rounded=True, rad=0.12)
        tx(slide, 10.85, 1.64, 1.5, 0.38, str(d["year"]), size=11, color=pal.on(pal["accent"]),
           bold=True, align=PP_ALIGN.CENTER)
    design_color = pal.tokens['on_primary_soft']   # token on_primary_soft
    pico_color = pal.tokens['on_primary']          # token on_primary
    tx(slide, 0.95, 2.06, 9.6, 0.3, d.get("design", "").upper(), size=9.5, color=design_color, bold=True)
    rows = [(l, d.get(f, "")) for l, f in (('P', 'population'), ('I', 'intervention'),
            ('C', 'comparison'), ('O', 'outcome')) if d.get(f)]
    rh, gp = 0.85, 0.12
    for i, (ltr, txt) in enumerate(rows):
        y = 2.65 + i * (rh + gp)
        rect(slide, SAFE_X0, y, 0.55, 0.55, fill=pal["primary"], line=None, rounded=True, rad=0.15)
        tx(slide, SAFE_X0, y, 0.55, 0.55, ltr, size=15, color=pico_color, bold=True,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
        tx(slide, 1.35, y, 6.5, rh, txt, size=11.0, color=pal["text"], line_sp=1.1)
    rect(slide, 8.15, 2.65, 4.5, 1.7, fill=pal["accent"], line=None, rounded=True)
    hero_fg = pal.on(pal["accent"])
    tx(slide, 8.3, 2.82, 4.2, 0.9, d.get("effect", ""), size=36, color=hero_fg,
       bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE)
    tx(slide, 8.35, 3.78, 4.1, 0.45, d.get("effect_desc", ""), size=10.5, color=hero_fg,
       align=PP_ALIGN.CENTER, line_sp=1.05)
    if d.get("limitation"):
        rect(slide, 8.15, 4.55, 4.5, 1.85, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, 8.15, 4.55, 0.06, 1.85, fill=pal["alert"], rounded=False)
        tx(slide, 8.4, 4.7, 4.0, 0.28, T('limitation'), size=9, color=pal["alert"], bold=True)
        tx(slide, 8.4, 5.0, 4.0, 1.3, d["limitation"], size=10.5, color=pal["text"], line_sp=1.1)
    footnote(slide, pal, d, y=6.60)  # credit = full citation (Vancouver)

def m_viva_question(slide, d, pal, info):
    """Viva voce: question card, model answer, examiner's note."""
    ink = pal.ink_accent()
    rect(slide, SAFE_X0, SAFE_Y0, 5.7, 4.95, fill=pal["primary"], line=None, rounded=True)
    tx(slide, 0.95, 1.75, 5.1, 0.3, T('viva'), size=10, color=pal.tokens['on_primary'], bold=True)
    if d.get("difficulty"):
        pill(slide, 0.95, 2.13, 1.7, 0.34, d["difficulty"].upper(), pal["accent"], pal.on(pal["accent"]))
    tx(slide, 0.95, 2.65, 5.1, 3.6, d.get("prompt", ""), size=16, color=pal.tokens['on_primary'],
       font=FONT_TITLE, line_sp=1.2)
    rect(slide, 6.55, SAFE_Y0, 6.1, 2.75, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    tx(slide, 6.8, 1.72, 5.6, 0.3, T('model_answer'), size=10, color=ink, bold=True)
    tx(slide, 6.8, 2.12, 5.6, 2.0, d.get("model_answer", ""), size=11.5, color=pal["text"], line_sp=1.12)
    rect(slide, 6.55, 4.5, 6.1, 2.0, fill=pal["soft"], line=pal["border"], lw=0.6, rounded=True)
    rect(slide, 6.55, 4.5, 0.06, 2.0, fill=pal["alert"], rounded=False)
    tx(slide, 6.8, 4.65, 5.6, 0.3, T('examiner'), size=10, color=pal["alert"], bold=True)
    tx(slide, 6.8, 5.02, 5.6, 1.35, d.get("examiner_note", ""), size=10.8, color=pal["text"], line_sp=1.1)
    footnote(slide, pal, d)

def m_mnemonic_card(slide, d, pal, info):
    """Acronym expansion: letter medallions + aligned terms + clinical use panel."""
    ink = pal.ink_accent()
    acr = d.get("acronym", "")
    exp = d.get("expansion", [])
    n = max(len(acr), 1)
    gap = 0.15
    cw = min(1.35, (CONTENT_W - gap * (n - 1)) / n)
    total = cw * n + gap * (n - 1)
    x0 = SAFE_X0 + (CONTENT_W - total) / 2.0
    if d.get("domain"):
        pill(slide, 10.30, 1.28, 2.35, 0.34, d["domain"].upper(), pal["soft"], pal["primary"])
    for i, ch in enumerate(acr):
        x = x0 + i * (cw + gap)
        rect(slide, x, 1.95, cw, 0.95, fill=pal["primary"], line=None, rounded=True, rad=0.12)
        tx(slide, x, 1.95, cw, 0.95, ch.upper(), size=28, color=pal.tokens['on_primary'], bold=True,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
        tx(slide, x - 0.075, 3.0, cw + 0.15, 0.8, exp[i] if i < len(exp) else "", size=9.5,
           color=pal["text"], align=PP_ALIGN.CENTER, line_sp=1.05)
    rect(slide, SAFE_X0, 4.1, CONTENT_W, 1.7, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    rect(slide, SAFE_X0, 4.1, CONTENT_W, 0.05, fill=pal["accent"], rounded=False)
    tx(slide, 0.95, 4.28, 11.4, 0.3, T('clinical_use'), size=10, color=ink, bold=True)
    tx(slide, 0.95, 4.62, 11.4, 1.05, d.get("usage", ""), size=12, color=pal["text"], line_sp=1.15)
    footnote(slide, pal, d)

def m_etymology(slide, d, pal, info):
    """Philological deconstruction: colored term fragments + root cards."""
    roots = d.get("roots", [])
    frags = d.get("fragments", [])
    ink = pal.ink_accent()
    tf = textframe(slide, SAFE_X0, 1.85, CONTENT_W, 1.25)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    frag_cols = [pal["primary"], ink]  # sin alert rojo decorativo: semantica falsa
    parts = frags if (frags and len(frags) == len(roots)) else [d.get("term", "")]
    for i, fr in enumerate(parts):
        r = p.add_run()
        r.text = fr
        r.font.size = Pt(40)
        r.font.bold = True
        r.font.name = FONT_TITLE
        r.font.color.rgb = C(frag_cols[i % len(frag_cols)] if len(parts) > 1 else pal["primary"])
    rect(slide, 5.87, 3.2, 1.6, 0.03, fill=pal["accent"], rounded=False)
    n = max(len(roots), 1)
    gap = 0.3
    w, sx = row_layout(SAFE_X0, CONTENT_W, n, gap)
    for i, rt in enumerate(roots):
        x = sx + i * (w + gap)
        rect(slide, x, 3.5, w, 1.85, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, x, 3.5, w, 0.05, fill=pal["accent"], rounded=False)
        tfc = textframe(slide, x + 0.2, 3.68, w - 0.4, SAFE_Y0)
        para(tfc, rt.get("part", ""), 15, pal["primary"], bold=True, font=FONT_TITLE, first=True,
             sp_after=3, align=PP_ALIGN.CENTER)
        para(tfc, T('origin') + " · " + rt.get("origin", "").upper(), 9, ink, bold=True,
             sp_after=5, align=PP_ALIGN.CENTER)
        para(tfc, rt.get("meaning", ""), 10.5, pal["text"], line_sp=1.1, align=PP_ALIGN.CENTER)
    rect(slide, SAFE_X0, 5.6, CONTENT_W, 0.8, fill=pal["soft"], line=pal["border"], lw=0.4, rounded=True)
    tx(slide, 0.9, 5.68, 11.5, 0.28, T('modern_use'), size=9.5, color=ink, bold=True)
    tx(slide, 0.9, 5.96, 11.5, 0.38, d.get("modern", ""), size=11.5, color=pal["primary"], bold=True)
    footnote(slide, pal, d)

def m_venn_diagram(slide, d, pal, info):
    """Set overlap (2-3 translucent circles); 'shared' = intersection zone (con halo)."""
    sets_l = d.get("sets", [])
    shared = d.get("shared", [])
    n = len(sets_l)
    def _circle(cx, cy, r, color):
        ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                                    Inches(2 * r), Inches(2 * r))
        ov.fill.solid(); ov.fill.fore_color.rgb = C(color)
        ov.line.color.rgb = C(pal["border"]); ov.line.width = Pt(1.25)
        ov.shadow.inherit = False
        set_fill_alpha(ov, 42)
    def _halo(cx, cy, w, h):
        shp = rect(slide, cx - w / 2 - 0.05, cy - h / 2 - 0.05, w + 0.10, h + 0.10,
                   fill=pal["card"], line=None, rounded=True, rad=0.10)
        set_fill_alpha(shp, 88)
    def _items(cx, cy, w, h, items, size=10.5):
        if not items:
            return
        tfi = textframe(slide, cx - w / 2, cy - h / 2, w, h, anchor=MSO_ANCHOR.MIDDLE)
        for i, it in enumerate(items):
            para(tfi, "• " + it, size, pal["text"], first=(i == 0), sp_after=4,
                 line_sp=1.08, align=PP_ALIGN.CENTER)
    def _label(cx, y, text):
        tx(slide, cx - 1.2, y, 2.4, 0.32, str(text).upper(), size=11, color=pal["primary"],
           bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE)
    if n == 3:
        for (cx, cyy), col in zip(((6.666, 3.0), (5.35, 4.72), (7.98, 4.72)),
                                  (pal["primary"], pal["accent"], pal["primary"])):
            _circle(cx, cyy, 1.5, col)
        _label(6.666, 1.5, sets_l[0].get("label", ""))
        _label(5.35, 6.28, sets_l[1].get("label", ""))
        _label(7.98, 6.28, sets_l[2].get("label", ""))
        # Zonas de items ampliadas y fuente menor: 3-4 items nao transbordam
        _items(6.666, 2.55, 1.9, 1.3, sets_l[0].get("items", []), size=9.5)
        _items(4.50, 5.30, 1.7, 1.5, sets_l[1].get("items", []), size=9.5)
        _items(8.83, 5.30, 1.7, 1.5, sets_l[2].get("items", []), size=9.5)
        _halo(6.666, 4.15, 1.35, 0.95)
        _items(6.666, 4.15, 1.35, 0.95, shared, size=9.0)
    else:
        _circle(5.55, 4.1, 1.85, pal["primary"])
        _circle(7.78, 4.1, 1.85, pal["accent"])
        _label(5.55, 1.95, sets_l[0].get("label", "") if sets_l else "")
        _label(7.78, 1.95, sets_l[1].get("label", "") if len(sets_l) > 1 else "")
        _items(4.6, 4.1, 1.8, 2.0, sets_l[0].get("items", []) if sets_l else [])
        _items(8.73, 4.1, 1.8, 2.0, sets_l[1].get("items", []) if len(sets_l) > 1 else [])
        _halo(6.665, 4.1, 1.45, 2.1)
        _items(6.665, 4.1, 1.45, 2.1, shared, size=9.5)
    footnote(slide, pal, d)

def m_references(slide, d, pal, info):
    """Vancouver-style numbered bibliography in two balanced columns."""
    refs = d.get("refs", [])
    ink = pal.ink_accent()
    n = len(refs)
    pill(slide, 10.85, 1.28, 1.8, 0.34, f"{T('references')} · {n}", pal["soft"], pal["primary"])
    per = max((n + 1) // 2, 1)
    for ci, col in enumerate([refs[:per], refs[per:]]):
        cx = SAFE_X0 + ci * 6.15
        for j, ref in enumerate(col):
            y = 2.05 + j * 0.92
            tx(slide, cx, y, 0.45, 0.35, f"{ci * per + j + 1}.", size=12, color=ink,
               bold=True, font=FONT_TITLE)
            tx(slide, cx + 0.5, y, 5.35, 0.82, ref, size=10.5, color=pal["text"], line_sp=1.1)
    footnote(slide, pal, d)

def m_authority_card(slide, d, pal, info):
    """Biografia con retrato enmarcado: nombre, fechas, bio y legado."""
    ink = pal.ink_accent()
    place_image(slide, d["image"], SAFE_X0 + 0.04, SAFE_Y0, 3.7, 4.95, pal)
    tx(slide, 4.75, 1.7, 7.8, 0.6, d.get("name", ""), size=24, color=pal["primary"],
       bold=True, font=FONT_TITLE)
    if d.get("dates"):
        tx(slide, 4.75, 2.38, 7.8, 0.3, d["dates"], size=11, color=ink, bold=True)
    if d.get("epithet"):
        tx(slide, 4.75, 2.72, 7.8, 0.35, d["epithet"], size=11.5, color=pal["text_muted"], italic=True)
    tx(slide, 4.75, 3.2, 7.8, 2.4, d.get("bio", ""), size=11.5, color=pal["text"], line_sp=1.18)
    if d.get("legacy"):
        rect(slide, 4.75, 5.75, 7.8, 0.75, fill=pal["soft"], line=pal["border"], lw=0.5, rounded=True)
        rect(slide, 4.75, 5.75, 0.06, 0.75, fill=pal["accent"], rounded=False)
        tx(slide, 5.0, 5.83, 7.3, 0.28, T('legacy'), size=9, color=ink, bold=True)
        tx(slide, 5.0, 6.1, 7.3, 0.35, d["legacy"], size=10.5, color=pal["primary"], bold=True)
    footnote(slide, pal, d)
