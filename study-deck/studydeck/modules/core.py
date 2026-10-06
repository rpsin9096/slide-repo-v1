# -*- coding: utf-8 -*-
"""Renderers tier core + classic: portada, division, definicion, flujo, tablas, metricas, protocolos, caso clinico y figura."""

import math
import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..contrast import ratio, shade_until, WCAG_AA_TEXT
from ..geometry import (
    CANVAS_H, CANVAS_W, CONTENT_W, SAFE_X0, SAFE_Y0,
    row_layout, stack_layout,
)
from ..i18n import T, bloom
from ..primitives import (
    C, bg, footnote, note_surface, para, pill, place_image, rect,
    set_fill_alpha, textframe, tx, unmeasured,
)
from ..typography import FONT_BODY, FONT_TITLE

def m_title(slide, d, pal, info):
    """Academic title slide with crest layout — supports optional author."""
    bg(slide, pal, pal["primary"])
    rect(slide, 0, 0, CANVAS_W, 0.18, fill=pal["accent"], rounded=False)
    rect(slide, 0, 7.32, CANVAS_W, 0.18, fill=pal["accent"], rounded=False)
    rect(slide, 0.8, 0.7, 11.733, 6.1, fill=None, line=pal["accent"], lw=1.0, rounded=True, rad=0.03)
    course_name = d.get("course", info.get("course", ""))
    title_txt = d.get("title") or course_name  # fallback: o curso como titulo
    course_color = pal.tokens['on_primary']
    inst_color = pal.tokens['on_primary_soft']
    sub_color = pal.subtitle()
    if course_name:
        tx(slide, 1.2, 1.35, 10.933, 0.4, course_name.upper(), size=12,
           color=course_color, bold=True, align=PP_ALIGN.CENTER, font=FONT_BODY)
    tx(slide, 1.2, 2.1, 10.933, 2.1, title_txt, size=38, color=pal.tokens['on_primary'],
       bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE, line_sp=1.08)
    rect(slide, 5.86, 4.45, 1.6, 0.03, fill=pal["accent"], rounded=False)
    tx(slide, 1.5, 4.75, 10.333, 0.7, d.get("subtitle", ""), size=15, color=sub_color,
       align=PP_ALIGN.CENTER, italic=True, font=FONT_BODY)
    inst = d.get("institution", info.get("institution", ""))
    if inst:
        tx(slide, 1.5, 5.85, 10.333, 0.4, inst.upper(), size=10, color=inst_color,
           align=PP_ALIGN.CENTER, font=FONT_BODY, bold=True)
    author = d.get("author", info.get("author", ""))
    if author:
        tx(slide, 1.5, 6.25, 10.333, 0.35, author, size=11, color=pal.tokens['on_primary'],
           align=PP_ALIGN.CENTER, font=FONT_BODY, italic=True)

def m_section_divider(slide, d, pal, info):
    """Divisor de secao — fondo primario de marca; autor opcional."""
    bg(slide, pal, pal["primary"])
    rect(slide, 0, 0, 0.4, 7.5, fill=pal["accent"], rounded=False)
    style = d.get("divider_style", T('chapter'))
    sub_color = pal.subtitle()
    if style:
        tx(slide, 1.2, 1.5, 11.0, 0.5, style.upper(), size=14, color=pal.tokens['on_primary'],
           bold=True, align=PP_ALIGN.LEFT, font=FONT_BODY)
    tx(slide, 1.2, 2.2, 11.0, 1.8, d.get("section_title", ""), size=36, color=pal.tokens['on_primary'],
       bold=True, align=PP_ALIGN.LEFT, font=FONT_TITLE, line_sp=1.1)
    rect(slide, 1.2, 4.25, 2.0, 0.03, fill=pal["accent"], rounded=False)
    tx(slide, 1.2, 4.6, 10.5, 1.5, d.get("section_subtitle", ""), size=15, color=sub_color,
       align=PP_ALIGN.LEFT, italic=True, line_sp=1.2, font=FONT_BODY)
    author = d.get("author", info.get("author", ""))
    if author:
        tx(slide, 1.2, 6.3, 10.5, 0.35, author, size=11, color=pal.tokens['on_primary_soft'],
           align=PP_ALIGN.LEFT, font=FONT_BODY, italic=True)

def m_definition(slide, d, pal, info):
    """Academic thesis/definition module with structured proposition cards."""
    img = bool(d.get("image"))
    cw = 7.6 if img else CONTENT_W
    rect(slide, SAFE_X0, SAFE_Y0, cw, 1.45, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    rect(slide, SAFE_X0, SAFE_Y0, 0.09, 1.45, fill=pal["primary"], rounded=False)
    pill(slide, 0.95, 1.68, 1.6, 0.24, T('central'), pal["soft"], pal["primary"], rad=0.08)
    tx(slide, 0.95, 1.98, cw - 0.5, 0.95, d.get("main_statement", ""), size=15.5,
       color=pal["primary"], bold=True, font=FONT_TITLE, line_sp=1.1)
    y = 3.15
    if d.get("elaboration"):
        rect(slide, SAFE_X0, y, cw, 0.85, fill=pal["soft"], line=pal["border"], lw=0.4, rounded=True)
        tx(slide, 0.85, y + 0.12, cw - 0.4, SAFE_X0, d["elaboration"], size=12.0,
           color=pal["text"], line_sp=1.15, font=FONT_BODY)
        y += 1.0
    kps = d.get("key_points", [])
    if kps:
        n = len(kps)
        gap = 0.22
        w = (cw - gap * (n - 1)) / n
        ch = 6.6 - y
        for i, kp in enumerate(kps):
            x = SAFE_X0 + i * (w + gap)
            rect(slide, x, y, w, ch, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
            rect(slide, x, y, w, 0.05, fill=pal["accent"], rounded=False)
            tf = textframe(slide, x + 0.2, y + 0.18, w - 0.4, ch - 0.3)
            para(tf, kp.get("title", ""), 12.5, pal["primary"], bold=True, font=FONT_TITLE, first=True, sp_after=4)
            para(tf, kp.get("desc", ""), 11.0, pal["text"], line_sp=1.1)
    if img:
        place_image(slide, d["image"], 8.45, SAFE_Y0, 4.2, 5.0, pal)
    footnote(slide, pal, d)

def m_flow_diagram(slide, d, pal, info):
    """Linear pathway; steps with decision:True render as diamonds (straight arrow = YES,
    bottom corridor bypass to step i+2 = NO, labeled by that step's `desc`)."""
    steps = d.get("steps", [])
    expl = d.get("explanation", "")
    yes_color = pal.ok_ink()
    sub_color = pal.subtitle()
    h = 3.6 if expl else 4.6
    y = 1.65
    n = max(len(steps), 1)
    aw = 0.35
    w = (CONTENT_W - aw * (n - 1)) / n
    yb = y + h + 0.12  # NO-branch corridor
    for i, st in enumerate(steps):
        x = SAFE_X0 + i * (w + aw)
        hi = st.get("highlight", False)
        dec = bool(st.get("decision"))
        if dec:
            dsz = min(w * 0.92, h * 0.72)
            dx, dy = x + (w - dsz) / 2.0, y + (h - dsz) / 2.0
            dm = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(dx), Inches(dy), Inches(dsz), Inches(dsz))
            dm.fill.solid(); dm.fill.fore_color.rgb = C(pal["card"])
            note_surface(pal["card"], dx, dy, dsz, dsz)
            dm.line.color.rgb = C(pal["accent"]); dm.line.width = Pt(1.2)
            dm.shadow.inherit = False
            tf = dm.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.04)
            para(tf, st.get("title", ""), 10.5, pal["primary"], bold=True, font=FONT_TITLE,
                 first=True, align=PP_ALIGN.CENTER, sp_after=0, line_sp=1.0)
        else:
            rect(slide, x, y, w, h, fill=(pal["primary"] if hi else pal["card"]),
                 line=(pal["accent"] if hi else pal["border"]), lw=1.0 if hi else 0.6, rounded=True)
            badge_bg = pal["accent"] if hi else pal["soft"]
            badge_fg = pal.on(pal["accent"]) if hi else pal["primary"]
            pill(slide, x + 0.18, y + 0.18, 0.8, 0.28, st.get("label", f"{T('phase')} {i+1}").upper(),
             badge_bg, badge_fg, size=9, rad=0.08)
            tf = textframe(slide, x + 0.2, y + 0.58, w - 0.4, h - 0.75)
            para(tf, st.get("title", ""), 13.5, pal.tokens['on_primary'] if hi else pal["primary"], bold=True,
                 font=FONT_TITLE, first=True, sp_after=6)
            para(tf, st.get("desc", ""), 11.0, sub_color if hi else pal["text"], line_sp=1.1)
        if i < n - 1:
            ax, ay = x + w + 0.05, y + (h / 2.0) - 0.12
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(ay), Inches(aw - 0.1), Inches(0.24))
            ar.fill.solid(); ar.fill.fore_color.rgb = C(pal["accent"])
            ar.line.fill.background(); ar.shadow.inherit = False
            bypass = dec and (i < n - 2) and not expl
            if dec:
                tx(slide, ax - 0.12, ay - 0.3, 0.6, 0.24, T("yes"), size=8.5,
                   color=yes_color, bold=True, align=PP_ALIGN.CENTER)
                if bypass:  # NO branch: elbow under the row to card i+2
                    xc = x + w / 2.0
                    xt = SAFE_X0 + (i + 2) * (w + aw) + w / 2.0
                    rect(slide, xc - 0.0125, y + h + 0.02, 0.025, yb - (y + h) - 0.02, fill=pal["alert"], rounded=False)
                    rect(slide, min(xc, xt), yb, abs(xt - xc), 0.025, fill=pal["alert"], rounded=False)
                    up = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(xt - 0.09), Inches(y + h + 0.04), Inches(0.18), Inches(0.18))
                    up.fill.solid(); up.fill.fore_color.rgb = C(pal["alert"])
                    up.line.fill.background(); up.shadow.inherit = False
                    # Caption DEBAJO del corredor: arriba colisionaba con la linea
                    tx(slide, (xc + xt) / 2.0 - 0.8, yb + 0.04, 1.6, 0.22,
                       st.get("desc") or T("no"), size=8.5, color=pal["alert"], bold=True, align=PP_ALIGN.CENTER)
                elif st.get("desc"):
                    # Sin bypass no hay corredor: la banda queda reservada al caption
                    tx(slide, x + 0.05, y + h + 0.18, w - 0.1, 0.22, st["desc"], size=9.0,
                       color=pal["text_muted"], align=PP_ALIGN.CENTER, line_sp=1.0)
        elif dec and st.get("desc"):
            # Decision as final step: no outgoing arrow — show NO-branch caption below.
            tx(slide, x + 0.05, y + h + 0.18, w - 0.1, 0.22, st["desc"], size=9.0,
               color=pal["text_muted"], align=PP_ALIGN.CENTER, line_sp=1.0)
    if expl:
        ey = y + h + 0.25
        rect(slide, SAFE_X0, ey, CONTENT_W, 6.6 - ey, fill=pal["soft"], line=pal["border"], lw=0.5, rounded=True)
        tx(slide, 0.9, ey + 0.1, 11.5, 6.4 - ey, expl, size=11.5, color=pal["text"],
           line_sp=1.12, anchor=MSO_ANCHOR.MIDDLE)
    footnote(slide, pal, d, y=6.66)

def _render_table_academic(slide, d, pal, y_top=1.6, header_key="headers", rows_key="rows"):
    """Academic table styling: crisp horizontal rules, soft alternating fills.
    Superficie marcada desconocida durante la tabla: el gate no mide celdas nativas."""
    headers = d.get(header_key, [])
    rows = d.get(rows_key, [])
    if not headers:
        raise KeyError(f"missing '{header_key}' (tabla con rows_key='{rows_key}')")
    with unmeasured():
        takeaway = d.get("takeaway", "")
        tbl_h = (4.1 if takeaway else 4.7)
        gf = slide.shapes.add_table(len(rows) + 1, len(headers),
                                    Inches(SAFE_X0), Inches(y_top), Inches(CONTENT_W), Inches(tbl_h))
        table = gf.table
        cw = d.get("col_widths")
        if cw and len(cw) == len(headers):
            total = sum(cw)
            for j, wgt in enumerate(cw):
                table.columns[j].width = Inches(CONTENT_W * wgt / total)
        for j, htxt in enumerate(headers):
            cell = table.cell(0, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C(pal["primary"])
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.margin_top = cell.margin_bottom = Inches(0.08)
            tf = cell.text_frame
            tf.word_wrap = True
            para(tf, str(htxt).upper(), 11.0, pal.tokens['on_primary'], bold=True, font=FONT_TITLE, first=True, sp_after=0)
        for i, row in enumerate(rows):
            is_even = (i % 2 == 0)
            row_bg = pal["card"] if is_even else pal["soft"]
            for j in range(len(headers)):
                val = row[j] if j < len(row) else ""
                cell = table.cell(i + 1, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = C(row_bg)
                cell.margin_left = cell.margin_right = Inches(0.14)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                tf = cell.text_frame
                tf.word_wrap = True
                is_lead = (j == 0)
                para(tf, str(val), 11.0, pal["primary"] if is_lead else pal["text"],
                     bold=is_lead, font=FONT_BODY, first=True, sp_after=0, line_sp=1.05)
    if takeaway:
        y = y_top + tbl_h + 0.2
        rect(slide, SAFE_X0, y, CONTENT_W, 0.58, fill=pal["soft"], line=pal["accent"], lw=0.9, rounded=True)
        tx(slide, 0.9, y + 0.05, 11.5, 0.48, T('takeaway') + takeaway,
           size=11.5, color=pal["primary"], bold=True, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)

def m_comparison_table(slide, d, pal, info):
    """Tabla comparativa de 3-4 opciones sobre 3-4 criterios."""
    _render_table_academic(slide, d, pal)
    footnote(slide, pal, d, y=6.58)

def m_synthesis_grid(slide, d, pal, info):
    """Malla densa de relaciones (matriz-resumo de una seccion)."""
    _render_table_academic(slide, d, pal, rows_key="matrix_data")
    footnote(slide, pal, d, y=6.58)

def m_stat_card(slide, d, pal, info):
    """Impactful hero metric slide with clinical narrative support."""
    rect(slide, SAFE_X0, SAFE_Y0, 4.4, 4.95, fill=pal["primary"], line=None, rounded=True)
    rect(slide, SAFE_X0, SAFE_Y0, 4.4, 0.06, fill=pal["accent"], rounded=False)
    tx(slide, 0.85, 2.2, 4.0, 1.6, d.get("stat_num", ""), size=52, color=pal.tokens['on_primary'],
       bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE)
    # Barra divisoria centrada respecto al bloque del dígito (x .85 w 4.0 → centro 2.85).
    rect(slide, 2.25, 4.0, 1.2, 0.025, fill=pal["accent"], rounded=False)
    tx(slide, 0.9, 4.25, 3.9, 1.8, d.get("stat_title", ""), size=13.5, color=pal.tokens['on_primary'],
       align=PP_ALIGN.CENTER, line_sp=1.15, font=FONT_BODY)
    img = bool(d.get("image"))
    bw = 3.6 if img else 7.3
    bx = 4.25 if img else 5.35
    blocks = d.get("narrative_blocks", [])
    n = max(len(blocks), 1)
    gap = 0.2
    bh = (4.95 - gap * (n - 1)) / n
    for i, b in enumerate(blocks):
        y = SAFE_Y0 + i * (bh + gap)
        rect(slide, bx, y, bw, bh, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, bx, y, 0.06, bh, fill=pal["accent"], rounded=False)
        tf = textframe(slide, bx + 0.25, y + 0.12, bw - 0.45, bh - 0.24, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, b.get("title", ""), 12.5, pal["primary"], bold=True, font=FONT_TITLE, first=True, sp_after=3)
        para(tf, b.get("desc", ""), 11.0, pal["text"], line_sp=1.1)
    if img:
        place_image(slide, d["image"], 8.05, SAFE_Y0, 4.6, 4.95, pal)
    footnote(slide, pal, d)

def m_checklist(slide, d, pal, info):
    """Clinical protocol verification checklist."""
    items = d.get("items", [])
    ok_color = pal.ok_ink()
    n = max(len(items), 1)
    gap = 0.16
    ih, y_first = stack_layout(SAFE_Y0, 4.95, n, gap, cap=1.0)
    rw = 7.5 if d.get("image") else CONTENT_W
    for i, it in enumerate(items):
        y = y_first + i * (ih + gap)
        checked = it.get("checked", True)
        mark, mcol = ("✓", ok_color) if checked else ("✕", pal["alert"])
        rect(slide, SAFE_X0, y, rw, ih, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, 0.85, y + (ih / 2.0) - 0.18, 0.36, 0.36, fill=pal["soft"], line=mcol, lw=1.2, rounded=True, rad=0.08)
        tx(slide, 0.85, y + (ih / 2.0) - 0.20, 0.36, 0.36, mark, size=14, color=mcol, bold=True, align=PP_ALIGN.CENTER)
        tf = textframe(slide, 1.4, y + 0.08, rw - 0.85, ih - 0.16, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, it.get("label", ""), 12.5, pal["primary"], bold=True, font=FONT_TITLE, first=True, sp_after=2)
        if it.get("desc"):
            para(tf, it["desc"], 10.8, pal["text"], line_sp=1.05)
    if d.get("image"):
        place_image(slide, d["image"], 8.35, SAFE_Y0, 4.3, 4.95, pal)
    footnote(slide, pal, d)

def m_fixed_schema_card(slide, d, pal, info):
    """Structured diagnostic criteria / entity schema."""
    fields = d.get("fields", [])
    n = max(len(fields), 1)
    gap = 0.16
    fh, y_first = stack_layout(SAFE_Y0, 4.95, n, gap, cap=1.1)
    for i, f in enumerate(fields):
        y = y_first + i * (fh + gap)
        rect(slide, SAFE_X0, y, 3.4, fh, fill=pal["primary"], line=None, rounded=True)
        tx(slide, 0.85, y, 3.0, fh, f.get("label", "").upper(), size=11.0, color=pal.on(pal["primary"]),
           bold=True, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
        rect(slide, 4.15, y, 8.5, fh, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        tx(slide, 4.35, y, 8.1, fh, f.get("value", ""), size=11.5, color=pal["text"],
           anchor=MSO_ANCHOR.MIDDLE, line_sp=1.1)
    footnote(slide, pal, d)

def _render_column_card(slide, pal, x, w, card, accent_color):
    rect(slide, x, SAFE_Y0, w, 4.95, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    rect(slide, x, SAFE_Y0, w, 0.48, fill=accent_color, rounded=False)
    tx(slide, x + 0.25, 1.62, w - 0.5, 0.35, card.get("header", "").upper(), size=10.5,
       color=pal.on(accent_color), bold=True, font=FONT_BODY)
    tx(slide, x + 0.25, 2.15, w - 0.5, 0.55, card.get("title", ""), size=15.0,
       color=pal["primary"], bold=True, font=FONT_TITLE)
    tf = textframe(slide, x + 0.25, 2.75, w - 0.5, 3.6)
    for i, pt in enumerate(card.get("points", [])):
        para(tf, "• " + pt, 11.2, pal["text"], first=(i == 0), sp_after=7, line_sp=1.1)

def m_side_by_side(slide, d, pal, info):
    """Dichotomous comparative module (e.g., Condition A vs Condition B)."""
    m_w = 5.88
    _render_column_card(slide, pal, SAFE_X0, m_w, d.get("left_card", {}), pal["primary"])
    _render_column_card(slide, pal, 6.77, m_w, d.get("right_card", {}), pal["accent"])
    footnote(slide, pal, d)

def m_taxonomy_tree(slide, d, pal, info):
    """Hierarchical classification tree."""
    cats = d.get("categories", [])
    ink = pal.ink_accent()
    n = max(len(cats), 1)
    gap = 0.28
    w, _x0 = row_layout(SAFE_X0, CONTENT_W, n, gap)
    root = d.get("root", d.get("title", ""))
    rect(slide, 4.15, 1.5, 5.0, 0.55, fill=pal["primary"], line=None, rounded=True)
    tx(slide, 4.25, 1.5, 4.8, 0.55, root.upper(), size=12.5, color=pal.tokens['on_primary'], bold=True,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
    rect(slide, 6.64, 2.05, 0.02, 0.35, fill=pal["accent"], rounded=False)
    rect(slide, SAFE_X0 + w / 2, 2.4, CONTENT_W - w, 0.02, fill=pal["accent"], rounded=False)
    for i, c in enumerate(cats):
        x = SAFE_X0 + i * (w + gap)
        rect(slide, x + w / 2 - 0.01, 2.4, 0.02, 0.3, fill=pal["accent"], rounded=False)
        rect(slide, x, 2.7, w, 3.85, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
        rect(slide, x, 2.7, w, 0.45, fill=pal["soft"], rounded=False)
        tx(slide, x + 0.15, 2.75, w - 0.3, 0.35, c.get("name", "").upper(), size=11.0,
           color=pal["primary"], bold=True, align=PP_ALIGN.CENTER, font=FONT_TITLE)
        tf = textframe(slide, x + 0.2, 3.25, w - 0.4, 3.2)
        para(tf, c.get("property", ""), 11.0, ink, italic=True, font=FONT_BODY, first=True, sp_after=6)
        for ex in c.get("examples", []):
            para(tf, "• " + ex, 10.5, pal["text"], sp_after=4, line_sp=1.08)
    footnote(slide, pal, d)

def m_algorithm(slide, d, pal, info):
    """Clinical algorithmic cascade with decision medallions."""
    steps = d.get("steps", [])
    n = max(len(steps), 1)
    gap = 0.16
    sh, y_first = stack_layout(SAFE_Y0, 4.95, n, gap, cap=1.2)
    for i, st in enumerate(steps):
        y = y_first + i * (sh + gap)
        hi = st.get("highlight", False)
        med_fill = pal["accent"] if hi else pal["primary"]
        rect(slide, SAFE_X0, y, 0.9, sh, fill=med_fill, line=None, rounded=True)
        tx(slide, SAFE_X0, y, 0.9, sh, st.get("label", f"{i+1:02d}"), size=16, color=pal.on(med_fill),
           bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_TITLE)
        if i < n - 1:
            rect(slide, 1.09, y + sh, 0.02, gap, fill=pal["border"], rounded=False)
        rect(slide, 1.7, y, 10.95, sh, fill=pal["card"],
             line=pal["accent"] if hi else pal["border"], lw=1.0 if hi else 0.6, rounded=True)
        tf = textframe(slide, 1.95, y + 0.08, 10.4, sh - 0.16, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, st.get("title", ""), 12.5, pal["primary"], bold=True, font=FONT_TITLE, first=True, sp_after=2)
        if st.get("desc"):
            para(tf, st["desc"], 10.8, pal["text"], line_sp=1.05)
    footnote(slide, pal, d)

def m_case_block(slide, d, pal, info):
    """NEJM-styled clinical vignette presentation card."""
    img = bool(d.get("image"))
    ink = pal.ink_accent()
    lw = 4.6 if img else 5.85
    rx = 5.45 if img else 6.75
    rw = 4.6 if img else 5.9
    tw = lw - 0.4
    rect(slide, SAFE_X0, SAFE_Y0, lw, 2.15, fill=pal["primary"], line=None, rounded=True)
    tx(slide, 0.85, 1.7, tw, 0.25, T('vignette'), size=9.5, color=pal.tokens['on_primary'],
       bold=True, font=FONT_BODY)
    tx(slide, 0.85, 1.98, tw, 1.6, d.get("case_stem", ""), size=11.5, color=pal.tokens['on_primary'], line_sp=1.1, font=FONT_BODY)
    rect(slide, SAFE_X0, 3.85, lw, 2.65, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    tf = textframe(slide, 0.85, 4.0, tw, 2.4)
    para(tf, T('findings'), 10.0, ink, bold=True, font=FONT_BODY, first=True, sp_after=5)
    for fdg in d.get("findings", []):
        para(tf, "• " + fdg, 10.5 if img else 11.0, pal["text"], sp_after=3, line_sp=1.05)
    rect(slide, rx, SAFE_Y0, rw, 2.15, fill=pal["card"], line=pal["border"], lw=0.6, rounded=True)
    tf_r = textframe(slide, rx + 0.25, 1.7, rw - 0.5, 1.9)
    para(tf_r, T('reasoning'), 10.0, ink, bold=True, font=FONT_BODY, first=True, sp_after=4)
    para(tf_r, d.get("reasoning", ""), 10.8, pal["text"], line_sp=1.1)
    rect(slide, rx, 3.85, rw, 1.15, fill=pal["soft"], line=pal["accent"], lw=0.9, rounded=True)
    tf_d = textframe(slide, rx + 0.25, 3.95, rw - 0.5, 0.95, anchor=MSO_ANCHOR.MIDDLE)
    para(tf_d, T('diagnosis'), 9.5, pal["alert"], bold=True, first=True, sp_after=2)
    para(tf_d, d.get("diagnosis", ""), 12.5, pal["primary"], bold=True, font=FONT_TITLE)
    rect(slide, rx, 5.15, rw, 1.35, fill=pal["accent"], line=None, rounded=True)
    pearl_fg = pal.on(pal["accent"])
    tf_p = textframe(slide, rx + 0.25, 5.25, rw - 0.5, 1.15, anchor=MSO_ANCHOR.MIDDLE)
    para(tf_p, T('pearl'), 9.5, pearl_fg, bold=True, first=True, sp_after=2)
    para(tf_p, d.get("clinical_pearl", ""), 10.8, pearl_fg, line_sp=1.08)
    if img:
        place_image(slide, d["image"], 10.25, SAFE_Y0, 2.4, 4.95, pal)
    footnote(slide, pal, d)

def m_figure(slide, d, pal, info):
    """Full-bleed clinical figure showcase."""
    place_image(slide, d["image"], SAFE_X0, SAFE_Y0, CONTENT_W, 4.75, pal)
    caption_txt = d.get("caption", "")
    if caption_txt:
        tx(slide, SAFE_X0, 6.30, CONTENT_W, 0.28, caption_txt, size=10.5, color=pal["text"], italic=True)
    footnote(slide, pal, d, y=6.62)
