# -*- coding: utf-8 -*-
"""Primitivas de dibujo: caja, texto, pastilla, pie e imagen.

Todas las primitivas pintan en un lienzo de python-pptx y dejan rastro de la
superficie visible para el gate de contraste. El gate se inyecta con
`set_gate()`: el render crea uno por deck y lo instala aqui.
"""

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

from . import i18n
from .geometry import CONTENT_W, FOOTER_Y, SAFE_X0
from .images import classify
from .palette import load_palette
from .typography import FONT_BODY, FONT_TITLE

_GATE = None
_IMG_CACHE = {}


def set_gate(gate):
    """set_gate: documented behavior of the module."""
    global _GATE
    _GATE = gate
    return gate


def _require_gate():
    if _GATE is None:
        raise RuntimeError('gate de contraste no instalado; llame set_gate() antes de render')
    return _GATE


def note_surface(fill, x, y, w, h):
    """note_surface: documented behavior of the module."""
    _require_gate().note_surface(fill, x, y, w, h)


def unmeasured():
    """unmeasured: documented behavior of the module."""
    return _require_gate().unmeasured()


def C(h):
    """C: documented behavior of the module."""
    return RGBColor.from_string(str(h))


def _apply_adjustment(shp, rad):
    """Aplica el radio de redondeo SI la forma lo admite.

    Devuelve el radio aplicado o None: la via de falla produce informacion
    (la forma nativa se queda como esta) en vez de tragarla.
    """
    try:
        shp.adjustments[0] = rad
        return rad
    except (ValueError, IndexError, AttributeError):
        return None


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.5, rounded=True, rad=0.04):
    """Caja con redondeo profesional suave y sin sombra heredada."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        _apply_adjustment(shp, rad)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C(fill)
        note_surface(fill, x, y, w, h)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = C(line)
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def tx(slide, x, y, w, h, text, size=13, color=None, bold=False, italic=False,
       align=None, font=None, anchor=None, sp_after=3, line_sp=1.08, deco=False):
    """Caja de texto simple; el par texto/superficie queda medido por el gate.

    Sin color crudo: el default resuelve via la paleta activa (capa semantica).
    """
    if font is None:
        font = FONT_BODY
    if color is None:
        color = load_palette()['text']
    if align is None:
        align = PP_ALIGN.LEFT
    if anchor is None:
        anchor = MSO_ANCHOR.TOP
    _require_gate().log_text(color, size, bold, deco, x + w / 2.0, y + h / 2.0)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split('\n') if isinstance(text, str) else text):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sp_after)
        p.line_spacing = line_sp
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = C(color)
        r.font.name = font
    return box


def para(tf, text, size, color, bold=False, italic=False, align=None,
         font=None, sp_before=0, sp_after=3, first=False, line_sp=1.08, deco=False):
    """Parrafo dentro de una caja existente (flujo vertical)."""
    if font is None:
        font = FONT_BODY
    if align is None:
        align = PP_ALIGN.LEFT
    box = getattr(tf, '_box', None)
    if box and _GATE is not None:
        _GATE.log_text(color, size, bold, deco, box[0] + box[2] / 2.0, box[1] + box[3] / 2.0)
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(sp_before)
    p.space_after = Pt(sp_after)
    p.line_spacing = line_sp
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = C(color)
    r.font.name = font
    return p


def _note_box(tf, x, y, w, h):
    """Registra la geometria de la caja para que el gate mida su texto.

    Devuelve la caja o None cuando el TextFrame no admite atributos: el
    `para()` de esa caja queda sin medicion de contraste, y el None lo dice.
    """
    try:
        tf._box = (float(x), float(y), float(w), float(h))
        return tf._box
    except AttributeError:
        return None


def textframe(slide, x, y, w, h, anchor=None):
    """textframe: documented behavior of the module."""
    if anchor is None:
        anchor = MSO_ANCHOR.TOP
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf._box = _note_box(tf, x, y, w, h)
    return tf


def pill(slide, x, y, w, h, text, fill, fg, size=9.0, font=None, rad=0.12):
    """Pastilla/badge: caja redondeada + texto centrado sobre la MISMA caja."""
    if font is None:
        font = FONT_BODY
    rect(slide, x, y, w, h, fill=fill, line=None, rounded=True, rad=rad)
    tx(slide, x, y, w, h, text, size=size, color=fg, bold=True,
       align=PP_ALIGN.CENTER, font=font)


def bg(slide, pal, color=None):
    """bg: documented behavior of the module."""
    rect(slide, 0, 0, 13.333, 7.5, fill=color or pal['bg'], line=None, rounded=False)


def header(slide, pal, data):
    """Cabecera editorial: kicker discreto + titular serif + regla de acento."""
    kicker = data.get('kicker', '')
    title = data.get('title', '')
    if kicker:
        tx(slide, SAFE_X0, 0.32, 11.0, 0.24, kicker.upper(), size=9.5,
           color=pal.ink_accent(), bold=True, font=FONT_BODY)
    tx(slide, SAFE_X0, 0.58, CONTENT_W, SAFE_X0, title, size=23,
       color=pal['primary'], bold=True, font=FONT_TITLE)
    rect(slide, SAFE_X0, 1.28, 1.3, 0.03, fill=pal['accent'], rounded=False)
    rect(slide, 1.98, 1.29, 10.7, 0.01, fill=pal['border'], rounded=False)


def footer(slide, pal, idx, course, config=None):
    """Pie discreto: curso (opcional) + numeracion (opcional)."""
    config = config or {}
    if config.get('SHOW_FOOTER', True) and course:
        tx(slide, SAFE_X0, FOOTER_Y, 8.5, 0.28, course.upper(), size=9.0,
           color=pal['text_muted'], font=FONT_BODY)
    if config.get('SHOW_SLIDE_NUMBER', True):
        tx(slide, 11.8, FOOTER_Y, 0.9, 0.28, f'{idx:02d}', size=9.0,
           color=pal['text_muted'], align=PP_ALIGN.RIGHT, font=FONT_TITLE, bold=True)


def footnote(slide, pal, data, y=6.62):
    """Pie analitico + credito editorial de la fuente."""
    fn = data.get('footnote', '')
    credit = data.get('credit', '')
    combined = fn
    if credit and not fn:
        combined = i18n.T('source').format(c=credit)
    elif credit and fn and credit not in fn:
        combined = i18n.T('join').format(f=fn, c=credit)
    if combined:
        tx(slide, SAFE_X0, y, CONTENT_W, 0.28, combined, size=9.0,
           color=pal['text_muted'], italic=True)


def image_size(path):
    """(w, h) en px; cada archivo se abre con PIL una sola vez por proceso."""
    if path not in _IMG_CACHE:
        from PIL import Image as PILImage
        with PILImage.open(path) as im:
            _IMG_CACHE[path] = im.size
    return _IMG_CACHE[path]


def _panel(slide, x, y, w, h, pal, label):
    """Marco reservado para una fuente que el motor NO puede descargar."""
    rect(slide, x, y, w, h, fill=pal['soft'], line=pal['border'], lw=0.75,
         rounded=True, rad=0.02)
    tx(slide, x + 0.2, y + h / 2 - 0.16, w - 0.4, 0.32, label, size=9.5,
       color=pal['text_muted'], align=PP_ALIGN.CENTER)


def place_image(slide, path, x, y, w, h, pal=None):
    """Coloca una imagen respecting aspect ratio; marco segun paleta.

    Solo los archivos locales se pintan como pixeles. Las fuentes web, AI y
    placeholder se dibujan como panel rotulado: el motor no descarga nada y no
    inventa imagenes.
    """
    scheme, value = classify(path)
    if scheme != 'local':
        label = {'web': 'FUENTE WEB · ' + value,
                 'ai': 'FIGURA AI · ' + value,
                 'placeholder': 'MARCADOR · ' + value}.get(scheme, str(value))
        _panel(slide, x, y, w, h, pal or _FALLBACK_PAL, label)
        return None
    iw, ih = image_size(value)
    ar = iw / float(ih)
    box_ar = w / float(h)
    if ar > box_ar:
        pw, ph = w, w / ar
    else:
        pw, ph = h * ar, h
    px = x + (w - pw) / 2.0
    py = y + (h - ph) / 2.0
    if pal:
        rect(slide, px - 0.04, py - 0.04, pw + 0.08, ph + 0.08, fill=pal['card'],
             line=pal['border'], lw=0.75, rounded=True, rad=0.02)
    return slide.shapes.add_picture(value, Inches(px), Inches(py), Inches(pw), Inches(ph))


_FALLBACK_PAL = None


def set_fallback_palette(pal):
    """set_fallback_palette: documented behavior of the module."""
    global _FALLBACK_PAL
    _FALLBACK_PAL = pal


def set_fill_alpha(shape, opacity_pct):
    """Transparencia sobre relleno solido (solapamiento de conjuntos)."""
    spPr = shape._element.spPr
    sF = spPr.find(qn('a:solidFill'))
    if sF is None:
        return
    clr = sF.find(qn('a:srgbClr'))
    if clr is None:
        return
    for old in clr.findall(qn('a:alpha')):
        clr.remove(old)
    etree.SubElement(clr, qn('a:alpha'), val=str(int(opacity_pct * 1000)))