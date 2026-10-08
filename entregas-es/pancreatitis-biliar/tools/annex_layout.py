# -*- coding: utf-8 -*-
"""Pós-processo da diapositiva 13 (anexo): compacta la tabla y añade el esquema didáctico.

Uso (desde la carpeta de la entrega):
    studydeck all manifest.py Pancreatitis_biliar_12_diapositivas_es.pptx
    python tools/annex_layout.py Pancreatitis_biliar_12_diapositivas_es.pptx

Ejecutar después de cada build: el motor sobrescribe el .pptx desde el manifest.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "..", "assets", "ecografia_vesicular_esquema.png")
ANNEX_INDEX = 12  # diapositiva 13 (base 0)
TABLE_TOP = 1.55
ROW_H = 0.40
IMG_TOP = 3.70
IMG_W = 12.0
IMG_H = 2.05  # 12 x 2.05 in, matches the figure size of diagram_ecografia.py


def main(path):
    prs = Presentation(path)
    slide = prs.slides[ANNEX_INDEX]
    tables = [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]
    if len(tables) != 1:
        raise SystemExit("La diapositiva 13 debe tener exactamente una tabla; encontradas: %d" % len(tables))
    frame = tables[0]
    frame.top = Inches(TABLE_TOP)
    for row in frame.table.rows:
        row.height = Inches(ROW_H)
    frame.height = Inches(ROW_H * len(frame.table.rows))

    pic = slide.shapes.add_picture(IMG, Inches(0.65), Inches(IMG_TOP), Inches(IMG_W), Inches(IMG_H))
    pic.name = "Esquema didáctico de ecografía vesicular"
    pic._element.nvPicPr.cNvPr.set(
        "descr",
        "Esquema didáctico de ecografía vesicular: bilis normal anecoica, cálculo hiperecogénico con "
        "sombra acústica posterior, barro biliar con nivel y ecos bajos, y pólipo fijo sin sombra. "
        "Ilustración, no imagen clínica.")
    prs.save(path)
    print("OK: diapositiva 13 actualizada en", path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python tools/annex_layout.py <archivo.pptx>")
    main(sys.argv[1])
